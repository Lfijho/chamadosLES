"""Gera DOCX e PDF a partir da fonte Markdown final, sem acessar a aplicação."""
from datetime import datetime, timezone
from pathlib import Path
import hashlib
import re
import struct
import subprocess
import tempfile
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path.cwd()
SOURCE = ROOT / "docs/entrega-sprint-1.md"
OUTPUT = SOURCE.with_suffix(".docx")


def element(tag, **attributes):
    node = OxmlElement(tag)
    for key, value in attributes.items():
        node.set(qn(key), str(value))
    return node


def field(paragraph, instruction):
    run = paragraph.add_run()
    run._r.append(element("w:fldChar", **{"w:fldCharType": "begin"}))
    text = element("w:instrText", **{"xml:space": "preserve"})
    text.text = f" {instruction} "
    run._r.append(text)
    run._r.append(element("w:fldChar", **{"w:fldCharType": "end"}))


def style_by_id(doc, style_id):
    return next(style for style in doc.styles if style.style_id == style_id)


def format_document(path):
    doc = Document(path)
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(1.65)
    section.left_margin = section.right_margin = Cm(1.7)
    section.header_distance = section.footer_distance = Cm(0.7)
    normal = doc.styles["Normal"]
    normal.font.name = "Liberation Sans"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string("202830")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12
    normal.paragraph_format.widow_control = True
    normal.element.get_or_add_rPr().append(element("w:lang", **{"w:val": "pt-BR"}))
    for name, size in [("Title", 20), ("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11.5)]:
        style = style_by_id(doc, name.replace(" ", ""))
        style.font.name = "Liberation Sans"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string("173A43")
        style.paragraph_format.space_before = Pt(7)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True
    doc.paragraphs[0].style = doc.styles["Title"]
    doc.paragraphs[0].paragraph_format.space_before = Pt(0)
    for name in ["BodyText", "FirstParagraph", "Compact"]:
        style = style_by_id(doc, name)
        style.paragraph_format.line_spacing = 1.12
        style.paragraph_format.space_after = Pt(3 if name == "Compact" else 6)
    for name in ["Caption", "Image Caption"]:
        if name in doc.styles:
            doc.styles[name].font.size = Pt(9)
            doc.styles[name].font.italic = False
            doc.styles[name].paragraph_format.space_before = Pt(4)
            doc.styles[name].paragraph_format.space_after = Pt(6)
    for name in ["Source Code", "Verbatim Char"]:
        if name in doc.styles:
            doc.styles[name].font.name = "Liberation Mono"
            doc.styles[name].font.size = Pt(9)
            if hasattr(doc.styles[name], "paragraph_format"):
                doc.styles[name].paragraph_format.space_after = Pt(3)
    for paragraph in list(doc.paragraphs):
        if paragraph.text.strip() == "DOCX_PAGEBREAK":
            following = paragraph._p.getnext()
            while following is not None and following.tag != qn("w:p"):
                following = following.getnext()
            assert following is not None
            properties = following.get_or_add_pPr()
            properties.append(element("w:pageBreakBefore"))
            paragraph._p.getparent().remove(paragraph._p)
        if paragraph._p.xpath(".//w:drawing"):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.keep_with_next = True
            paragraph.paragraph_format.space_after = Pt(0)
        if paragraph.style.name in ["Caption", "Image Caption"]:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    # Pandoc recebe limites de largura/altura; a figura mantém a proporção original.
    for shape in doc.inline_shapes:
        relationship = shape._inline.graphic.graphicData.pic.blipFill.blip.embed
        blob = doc.part.related_parts[relationship].blob
        if blob.startswith(b"\x89PNG\r\n\x1a\n"):
            width, height = struct.unpack(">II", blob[16:24])
            scale = min(shape.width / width, shape.height / height)
            shape.width, shape.height = int(width * scale), int(height * scale)
    for table in doc.tables:
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        headers = [cell.text for cell in table.rows[0].cells]
        if len(headers) == 2:
            ratios = [0.11, 0.89] if headers[0] == "ID" else [0.27, 0.73]
        elif headers[0] == "Integrante":
            ratios = [0.29, 0.22, 0.49]
        elif headers[-1] == "Sprint":
            ratios = [0.12, 0.73, 0.15]
        elif headers[0] == "ID":
            ratios = [0.18, 0.47, 0.35]
        else:
            ratios = [0.27, 0.73]
        width = section.page_width - section.left_margin - section.right_margin
        for index, column in enumerate(table.columns):
            column.width = int(width * ratios[index])
        props = table._tbl.tblPr
        borders = element("w:tblBorders")
        for side in ["top", "left", "bottom", "right", "insideH", "insideV"]:
            borders.append(element("w:" + side, **{"w:val": "single", "w:sz": "4", "w:color": "CCD3D7"}))
        props.append(borders)
        for row_index, row in enumerate(table.rows):
            row._tr.get_or_add_trPr().append(element("w:cantSplit"))
            if row_index == 0:
                row._tr.get_or_add_trPr().append(element("w:tblHeader"))
            for cell_index, cell in enumerate(row.cells):
                cell.width = int(width * ratios[cell_index])
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                if row_index == 0:
                    cell._tc.get_or_add_tcPr().append(element("w:shd", **{"w:fill": "EDF1F2"}))
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_before = Pt(3)
                    paragraph.paragraph_format.space_after = Pt(3)
                    paragraph.paragraph_format.line_spacing = 1.05
                    for run in paragraph.runs:
                        run.font.size = Pt(9.5)
                        if row_index == 0:
                            run.font.bold = True
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Página ")
    field(footer, "PAGE")
    footer.add_run(" de ")
    field(footer, "NUMPAGES")
    for run in footer.runs:
        run.font.size = Pt(9)
    doc.settings.element.append(element("w:updateFields", **{"w:val": "true"}))
    doc.core_properties.author = "Leandro Pinzon Filho"
    doc.core_properties.last_modified_by = "Leandro Pinzon Filho"
    doc.core_properties.title = "HelpDesk: entrega da Sprint 1"
    doc.core_properties.subject = "Laboratório de Engenharia de Software"
    doc.core_properties.created = doc.core_properties.modified = datetime.now(timezone.utc)
    doc.save(path)


def main():
    source = SOURCE.read_text()
    assert len(re.findall(r"^## \d+\. ", source, re.M)) == 10
    images = re.findall(r"!\[[^\]]*\]\(([^)]+)\)", source)
    assert len(images) == 5  # Uma captura e quatro desenhos UML.
    assert all((SOURCE.parent / image).is_file() for image in images)
    with tempfile.TemporaryDirectory(prefix="helpdesk-docs-") as temporary:
        draft = Path(temporary) / "delivery.md"
        draft.write_text(source.replace("<!-- pagebreak -->", "DOCX_PAGEBREAK"))
        subprocess.run(["pandoc", str(draft), "--from=markdown+link_attributes-smart", "--to=docx",
                        "--no-highlight", "--resource-path=" + str(SOURCE.parent), "--output=" + str(OUTPUT)], check=True)
        format_document(OUTPUT)
        subprocess.run(["libreoffice", "-env:UserInstallation=file://" + temporary + "/lo-profile",
                        "--headless", "--convert-to", "pdf:writer_pdf_Export", "--outdir",
                        str(SOURCE.parent), str(OUTPUT)], check=True)
    pdf = SOURCE.with_suffix(".pdf")
    assert pdf.is_file()
    info = subprocess.check_output(["pdfinfo", str(pdf)], text=True)
    pages = re.search(r"^Pages:\s+(\d+)", info, re.M).group(1)
    print(f"Fonte SHA-256: {hashlib.sha256(source.encode()).hexdigest()}")
    print(f"DOCX e PDF gerados da mesma fonte. Páginas: {pages}")
    print(subprocess.check_output(["pandoc", "--version"], text=True).splitlines()[0])
    print(subprocess.check_output(["libreoffice", "--version"], text=True).strip())


if __name__ == "__main__":
    main()
